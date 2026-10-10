import json
from pathlib import Path
from collections import Counter
import openpyxl

root = Path(__file__).resolve().parents[1]
out = root / 'outputs/word_library_20261010'
spec = json.loads((out / 'workbook_spec.json').read_text(encoding='utf-8'))
data = json.loads((root / 'analysis/word_library_workbook_data.json').read_text(encoding='utf-8'))
book = openpyxl.load_workbook(spec['outputPath'], data_only=False)
cache = openpyxl.load_workbook(spec['outputPath'], data_only=True)
expected = {
    'Observation': data['library']['observation'],
    'Recommendation': data['library']['recommendation'],
    '待确认合并': data['review'],
    '确认项明细': data['review_details'],
    '原始记录': data['raw'],
    '来源报告': data['sources'],
}
checked = 0
for name, rows in expected.items():
    sheet = book[name]
    assert sheet.freeze_panes == 'C7', (name, sheet.freeze_panes)
    assert len(sheet.tables) == 1, name
    table = list(sheet.tables.values())[0]
    assert table.ref.endswith(str(len(rows)+6)), (name, table.ref)
    assert table.autoFilter is not None, name
    for rownum, row in enumerate(rows, 7):
        for colnum, value in enumerate(row, 1):
            actual = sheet.cell(rownum, colnum).value
            expected_value = "'" + value if isinstance(value, str) and value.startswith('=') else value
            assert (actual or '') == (expected_value or ''), (name, rownum, colnum)
            checked += 1
    assert all((dim.height or 0) <= 409 for dim in sheet.row_dimensions.values()), name

validation_checks = [('Observation', 'B7:B628'), ('Recommendation', 'B7:B181'), ('待确认合并', 'C7:C309')]
for name, target in validation_checks:
    rules = list(book[name].data_validations.dataValidation)
    assert len(rules) == 1 and str(rules[0].sqref) == target and rules[0].type == 'list', (name, rules)

assert all(book['待确认合并'].cell(r, 3).value == '待确认' for r in range(7, 310))
assert all(book['待确认合并'].cell(r, 13).data_type == 'f' for r in range(7, 310))
assert all(cache['待确认合并'].cell(r, 13).value == '待确认' for r in range(7, 310))
assert sum(cache['Observation'].cell(r, 5).value for r in range(7, 629)) == 791
assert sum(cache['Recommendation'].cell(r, 5).value for r in range(7, 182)) == 505
errors = [(sh.title, cell.coordinate, cell.value) for sh in cache for row in sh for cell in row if cell.data_type == 'e']
assert not errors, errors
report = {'verified': True, 'exact_source_cells_checked': checked,
          'rows': {name: len(rows) for name, rows in expected.items()},
          'native_tables': 6, 'dropdowns': 3, 'review_status_formulas': 303,
          'original_source_records': dict(Counter(r[6] for r in data['raw'])),
          'formula_errors': errors}
(out / 'native_verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=True))
