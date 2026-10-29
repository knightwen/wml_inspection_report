"""Exercise actual desktop Excel conversion, print pagination and PDF rendering."""
from pathlib import Path
from io import BytesIO
import sys,secrets,threading,json,re
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'workbench'))
from PIL import Image
from app import Store,RESOURCES
from preview import convert_excel_pdf,render_pdf

def main():
    folder=Path(__file__).resolve().parents[1]/'.build/qa/calculation-native'
    store=Store(folder)
    project=store.new('calculation')
    project['cal_cover']['title']='STRUCTURAL DESIGN REPORT'
    project['cal_cover']['modules'][0].update(title='PROJECT OVERVIEW',blocks=[{'id':secrets.token_hex(16),'type':'text','text':'Editable cover content verified in Excel and PDF.'}])
    project['metadata'].update(site='CAL WORKBENCH QA',project_number='QA',report_number='QA-CAL-001',location='Test location',subject='Printable report integration',report_date='2026-10-28',prepared_by='QA Author')
    text='External analysis output. This workbook compiles the report and does not execute engineering calculations. '
    project['modules'][0]['blocks']=[{'id':secrets.token_hex(16),'type':'text','text':text*100}, {'id':secrets.token_hex(16),'type':'table','rows':[['Case','Result'],['=1+1','Text only'],['QA','OK']]}]
    out=BytesIO();Image.new('RGB',(1000,600),'#8caf74').save(out,'PNG')
    photo=store.upload(out.getvalue(),'qa-figure.png');photo['caption']='Figure QA — imported analysis result'
    project['cal_cover']['modules'][0]['blocks'].extend([{'id':secrets.token_hex(16),'type':'table','rows':[['Cover input','Value'],['External result','=2+2']]},{'id':secrets.token_hex(16),'type':'image','photo':{**photo,'caption':'Cover figure QA'},'width':35}])
    project['modules'][1]['page_break']=True
    project['modules'][1]['blocks']=[{'id':secrets.token_hex(16),'type':'image','photo':photo,'width':80}]
    snapshot=store.save(project)
    result=store.export_snapshot(snapshot,'xlsx')
    store.export_snapshot(snapshot,'docx',export_id=result['id'])
    source=store.path/'exports'/(result['id']+'.xlsx')
    cancel=threading.Event();pdf=convert_excel_pdf(source,RESOURCES,cancel)
    pages=render_pdf(pdf,source,lambda done,total:None,cancel)
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(str(pdf)) as doc:
        texts=[]
        for page in doc:
            tp=page.get_textpage();texts.append(tp.get_text_range());tp.close();page.close()
    assert pages>=4,(pages,texts)
    assert all('STRUCTURAL DESIGN REPORT' in t and 'QA-CAL-001' in t for t in texts),texts
    assert 'PROJECT OVERVIEW' in texts[0] and 'Editable cover content' in texts[0]
    with pdfium.PdfDocument(str(pdf)) as doc:
        assert all(any(obj.type==3 for obj in page.get_objects()) for page in doc),'Missing repeated worksheet logo'
    assert '=1+1' in '\n'.join(texts)
    assert '=2+2' in '\n'.join(texts) and 'Cover figure QA' in '\n'.join(texts)
    from openpyxl import load_workbook
    book=load_workbook(source,read_only=True,data_only=True)
    assert book['Cover']['K1'].value=='STRUCTURAL DESIGN REPORT'
    normal=lambda text:re.sub(r'\s+','',str(text)).casefold()
    for _,contents_row,title in list(book['_Outline'].values)[1:]:
        actual=next(i+1 for i,text in enumerate(texts[2:],2) if 'CONTENTS' not in text and normal(title) in {normal(line) for line in text.splitlines()})
        assert int(book['Contents'].cell(contents_row,19).value)==actual,(title,actual)
    book.close()
    assert (source.with_suffix('.docx')).exists()
    summary={'pages':pages,'source':str(source),'pdf':str(pdf),'word':str(source.with_suffix('.docx'))}
    (folder/'result.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    print(json.dumps(summary))
    store.previews.close()

if __name__=='__main__':main()
