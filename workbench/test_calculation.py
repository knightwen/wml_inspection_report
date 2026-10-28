"""CAL data, recovery, exports and snapshot regression tests."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import json
import secrets
import tempfile
import threading
import unittest
import zipfile
from PIL import Image
from openpyxl import load_workbook
from lxml import etree as E
from app import Store, RESOURCES, RevisionConflict
from report_types import walk_modules
from preview import PreviewManager


class CalculationTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.store=Store(self.tmp.name)
        self.project=self.store.new('calculation')

    def tearDown(self):
        self.store.previews.close()
        self.tmp.cleanup()

    def photo(self):
        out=BytesIO();Image.new('RGB',(800,500),'#92d050').save(out,'PNG')
        return self.store.upload(out.getvalue(),'test.png')

    def content(self):
        module=self.project['modules'][0]
        photo=self.photo();photo['caption']='Test figure'
        module['blocks']=[{'id':secrets.token_hex(16),'type':'text','text':'External calculation result.\nSecond paragraph.'},
          {'id':secrets.token_hex(16),'type':'table','rows':[['Item','Result'],['=1+1','OK']]},
          {'id':secrets.token_hex(16),'type':'image','photo':photo,'width':70}]
        return module

    def test_blank_defaults_and_independent_report_structures(self):
        self.assertTrue(all(v=='' for v in self.project['metadata'].values()))
        self.assertEqual(len(list(walk_modules(self.project['modules']))),9)
        self.assertTrue(all(not m['blocks'] for _,m in walk_modules(self.project['modules'])))
        self.project['modules'][0]['title']='My chapter'
        self.assertEqual(self.store.new('calculation')['modules'][0]['title'],'References')
        self.assertEqual(self.store.new()['template_id'],'inspection')
        with self.assertRaises(ValueError):self.store.new('unknown')

    def test_save_conflict_and_template_identity(self):
        self.content();stale=deepcopy(self.project)
        saved=self.store.save(self.project)
        self.assertEqual(self.store.load(saved['id'])['modules'],saved['modules'])
        self.assertEqual(self.store.list()[0]['template_id'],'calculation')
        with self.assertRaises(RevisionConflict):self.store.save(stale)
        saved['template_id']='inspection';saved['metadata']['report_date']='2026-10-28'
        with self.assertRaises(ValueError):self.store.save(saved)

    def test_backup_restore_and_recovery_remap_every_image(self):
        module=self.content();original=module['blocks'][-1]['photo']['id']
        # Unsaved in-memory changes are included in recovery backups.
        module['title']='Unsent title'
        restored=self.store.restore(self.store.backup_snapshot(self.project))
        image=restored['modules'][0]['blocks'][-1]['photo']
        self.assertNotEqual(original,image['id'])
        self.assertEqual(self.store.photo_path(original).read_bytes(),self.store.photo_path(image['id']).read_bytes())
        self.assertEqual(restored['modules'][0]['title'],'Unsent title')
        self.assertNotEqual(self.project['id'],restored['id'])
        self.store.validate(restored)

    def test_invalid_structure_rejected_without_overwriting(self):
        good=deepcopy(self.project)
        bad=[]
        p=deepcopy(good);p['modules'][0]['title']='';bad.append(p)
        p=deepcopy(good);p['modules'][0]['id']=p['modules'][1]['id'];bad.append(p)
        p=deepcopy(good);p['modules'][0]['blocks']=[{'id':secrets.token_hex(16),'type':'table','rows':[['A'],['B','C']]}];bad.append(p)
        p=deepcopy(good);p['modules'][0]['blocks']=[{'id':secrets.token_hex(16),'type':'image','photo':{'id':secrets.token_hex(16)},'width':100}];bad.append(p)
        p=deepcopy(good);node=p['modules'][0]
        for _ in range(7):node['children']=[deepcopy(good['modules'][0])];node=node['children'][0];node['id']=secrets.token_hex(16)
        bad.append(p)
        for p in bad:
            with self.assertRaises(ValueError):self.store.save(p)
        self.assertEqual(self.store.load(good['id']),good)

    def test_xlsx_editable_cells_no_formulas_macros_or_historical_content(self):
        self.content();result=self.store.export_snapshot(self.project,'xlsx')
        path=self.store.path/'exports'/(result['id']+'.xlsx')
        book=load_workbook(path)
        self.assertEqual(book.sheetnames,['Cover','Contents','Report','_Outline'])
        cells=[c for s in book for row in s for c in row if c.value is not None]
        self.assertTrue(any(c.value=='=1+1' and c.data_type=='s' for c in cells))
        self.assertFalse(any(c.data_type=='f' for c in cells))
        self.assertEqual(book['Report'].print_title_rows,'$1:$9')
        self.assertEqual(book['_Outline'].sheet_state,'veryHidden')
        self.assertEqual(len(book['Report']._images),2)
        for name in ['Cover','Contents','Report']:
            self.assertEqual(book[name]['F1'].value,'CALCULATION REPORT')
            self.assertEqual(book[name].sheet_view.view,'normal')
            self.assertTrue(book[name]._images)
        with zipfile.ZipFile(path) as z:
            self.assertFalse(any('vba' in n.lower() for n in z.namelist()))
            self.assertTrue(any(n.startswith('xl/media/') for n in z.namelist()))
            all_xml=b''.join(z.read(n) for n in z.namelist() if n.endswith('.xml'))
            self.assertNotIn(b'Magnet Support',all_xml)
            for i in range(1,4):self.assertIn(b'CALCULATION REPORT',z.read(f'xl/worksheets/sheet{i}.xml'))

    def test_custom_cover_titles_content_and_borders_survive_exports_and_backup(self):
        self.content()
        self.project['cal_cover'].update(title='DESIGN & CALCULATION REPORT',introduction_title='PURPOSE',materials_title='DESIGN MATERIALS',introduction='Editable cover introduction',materials='Editable material notes',revisions=[{'revision':'B','date':'28/10/2026','status':'Review','designed':'QA','checked':'QC','approved':'AP'}])
        saved=self.store.save(self.project)
        restored=self.store.restore(self.store.backup(saved['id']))
        self.assertEqual(restored['cal_cover'],saved['cal_cover'])
        result=self.store.export_snapshot(restored,'xlsx')
        book=load_workbook(self.store.path/'exports'/(result['id']+'.xlsx'))
        for name in ['Cover','Contents','Report']:
            self.assertEqual(book[name]['F1'].value,'DESIGN & CALCULATION REPORT')
            self.assertEqual([getattr(book[name]['A5'].border,edge).style for edge in ['left','right','top','bottom']],['thin']*4)
        cover_values=[c.value for row in book['Cover'] for c in row]
        for text in ['PURPOSE','DESIGN MATERIALS','Editable cover introduction','Editable material notes','Review']:
            self.assertIn(text,cover_values)
        cells=[c for row in book['Report'] for c in row if c.value=='=1+1']
        self.assertEqual(len(cells),1)
        self.assertEqual([getattr(cells[0].border,edge).style for edge in ['left','right','top','bottom']],['thin']*4)
        revision=next(c for row in book['Cover'] for c in row if c.value=='REVISION')
        self.assertEqual([getattr(revision.border,edge).style for edge in ['left','right','top','bottom']],['thin']*4)
        result=self.store.export_snapshot(restored,'docx')
        with zipfile.ZipFile(self.store.path/'exports'/(result['id']+'.docx')) as z:
            self.assertIn('DESIGN & CALCULATION REPORT',''.join(E.fromstring(z.read('word/header1.xml')).itertext()))
            for text in ['PURPOSE','DESIGN MATERIALS','Editable cover introduction']:
                self.assertIn(text,''.join(E.fromstring(z.read('word/document.xml')).itertext()))

    def test_legacy_cover_without_title_keeps_default_branding(self):
        for key in ['title','introduction_title','materials_title']:self.project['cal_cover'].pop(key)
        self.store.save(self.project)
        result=self.store.export_snapshot(self.project,'xlsx')
        book=load_workbook(self.store.path/'exports'/(result['id']+'.xlsx'))
        self.assertEqual(book['Cover']['F1'].value,'CALCULATION REPORT')

    def test_docx_content_and_repeated_header(self):
        self.content();result=self.store.export_snapshot(self.project,'docx')
        path=self.store.path/'exports'/(result['id']+'.docx')
        with zipfile.ZipFile(path) as z:
            doc=E.fromstring(z.read('word/document.xml'))
            content=''.join(doc.itertext())
            self.assertIn('External calculation result.',content)
            self.assertIn('=1+1',content)
            self.assertIn('Test figure',content)
            self.assertIn(b'CALCULATION REPORT',z.read('word/header1.xml'))
            self.assertTrue(any(n.startswith('word/media/') for n in z.namelist()))
        with self.assertRaises(ValueError):self.store.export_snapshot(self.project,'pdf')

    def test_long_multilingual_table_cells_are_not_truncated(self):
        content='\u6c49'*1996+'TAIL'
        self.project['modules'][0]['blocks']=[{'id':secrets.token_hex(16),'type':'table','rows':[[content]+['']*19]}]
        result=self.store.export_snapshot(self.project,'xlsx')
        book=load_workbook(self.store.path/'exports'/(result['id']+'.xlsx'))
        text=''.join(str(c.value or '') for row in book['Report'] for c in row)
        self.assertEqual(text.count('\u6c49'),1996)
        self.assertIn('TAIL',text.replace('\n',''))
        self.assertLessEqual(max(r.height or 15 for r in book['Report'].row_dimensions.values()),240)

    def test_excel_preview_snapshot_does_not_block_save(self):
        entered=threading.Event();release=threading.Event()
        def convert(source,resources,cancel):
            entered.set();release.wait(5);pdf=source.with_suffix('.pdf');pdf.write_bytes(b'%PDF-test');return pdf
        previews=PreviewManager(self.store,RESOURCES,excel_converter=convert,excel_check=lambda:{'available':True},renderer=lambda *args:1)
        self.store.previews.close();self.store.previews=previews
        job=previews.start(self.store.report_snapshot(self.project['id']))
        try:
            self.assertTrue(entered.wait(5))
            self.project['modules'][0]['title']='After snapshot'
            self.store.save(self.project)
        finally:release.set()
        for thread in previews.threads:thread.join(5)
        result=previews.get(job['id']);self.assertEqual(result['status'],'ready')
        self.assertEqual(result['format'],'xlsx')
        book=load_workbook(self.store.path/'exports'/(result['report_id']+'.xlsx'))
        self.assertIn('1  References',[c.value for row in book['Report'] for c in row])
        self.assertNotIn('1  After snapshot',[c.value for row in book['Report'] for c in row])
        self.assertTrue((self.store.path/'exports'/(result['report_id']+'.docx')).exists())


if __name__=='__main__':unittest.main(verbosity=2)
